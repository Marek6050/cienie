"""Składa PPTX z zrzutów out/prezentacja/slajd-NN.png i notatek z notatki.json.

    node scripts/eksport-prezentacji.mjs && python3 scripts/zloz-pptx.py
Wymaga: pip install python-pptx
"""
import glob
import json

from pptx import Presentation
from pptx.util import Inches

WYJSCIE = "Cienie-Rzeczypospolitej-prezentacja-finalna.pptx"

notatki = json.load(open("components/prezentacja/notatki.json", encoding="utf-8"))
zdjecia = sorted(glob.glob("out/prezentacja/slajd-*.png"))
assert len(zdjecia) == len(notatki), f"{len(zdjecia)} zrzutów, {len(notatki)} notatek"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
pusty = prs.slide_layouts[6]

for zdjecie, n in zip(zdjecia, notatki):
    slajd = prs.slides.add_slide(pusty)
    slajd.shapes.add_picture(zdjecie, 0, 0, prs.slide_width, prs.slide_height)
    tekst = f"Mówi: {n['kto']}\nCzas: {n['czas']}\nKryterium: {n['kryterium']}\n\n" + "\n".join(f"- {p}" for p in n["punkty"])
    slajd.notes_slide.notes_text_frame.text = tekst

prs.save(WYJSCIE)
print("zapisano", WYJSCIE, len(zdjecia), "slajdów")
